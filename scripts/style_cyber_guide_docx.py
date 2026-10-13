#!/usr/bin/env python3
"""Apply Remnant Security styling to the Pandoc-generated editable guide."""

from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "081524"
CYAN = "00A98F"
AMBER = "D97706"
LIGHT_CYAN = "DDFBF5"
LIGHT_AMBER = "FFF3D8"
MUTED = "557086"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def keep_table_row_together(row) -> None:
    """Prevent a comparison row from splitting across two pages."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def move_paragraph_before(paragraph, before_paragraph) -> None:
    before_paragraph._p.addprevious(paragraph._p)


def move_paragraph_after(paragraph, after_paragraph) -> None:
    after_paragraph._p.addnext(paragraph._p)


def set_picture_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def style_document(input_path: Path, output_path: Path, logo_path: Path, qr_path: Path) -> None:
    doc = Document(input_path)

    for section in doc.sections:
        section.top_margin = Inches(0.68)
        section.bottom_margin = Inches(0.68)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        section.header_distance = Inches(0.28)
        section.footer_distance = Inches(0.28)

        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = "REMNANT SECURITY  /  HOW TO GET INTO CYBERSECURITY"
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in hp.runs:
            run.font.name = "Avenir Next"
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(CYAN)

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run("remnantsecurity.com/#waitlist   |   ")
        run.font.name = "Avenir Next"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
        add_page_field(fp)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Avenir Next"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    heading_specs = {
        "Title": (32, NAVY),
        "Heading 1": (30, NAVY),
        "Heading 2": (22, CYAN),
        "Heading 3": (14, AMBER),
    }
    for name, (size, color) in heading_specs.items():
        style = styles[name]
        style.font.name = "Avenir Next"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True

    for paragraph in doc.paragraphs:
        if (
            paragraph.style.name == "Heading 2"
            and paragraph.text != "A practical field guide for hungry beginners"
        ):
            paragraph.paragraph_format.page_break_before = True
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.style.name == "Block Text":
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.right_indent = Inches(0.15)
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(8)
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string(CYAN)

    first_paragraph = doc.paragraphs[0]
    first_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in first_paragraph.runs:
        run.font.size = Pt(32)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(NAVY)

    if len(doc.paragraphs) > 1:
        doc.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in doc.paragraphs[1].runs:
            run.font.color.rgb = RGBColor.from_string(AMBER)

    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_shape = logo_p.add_run().add_picture(str(logo_path), width=Inches(0.72))
    set_picture_alt_text(
        logo_shape,
        "Remnant Security logo",
        "Shield-shaped Remnant Security logo with a padlock in the center.",
    )
    move_paragraph_before(logo_p, first_paragraph)

    cover_anchor = doc.paragraphs[min(5, len(doc.paragraphs) - 1)]
    qr_p = doc.add_paragraph()
    qr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qr_shape = qr_p.add_run().add_picture(str(qr_path), width=Inches(0.68))
    set_picture_alt_text(
        qr_shape,
        "Remnant Security waitlist QR code",
        "QR code linking to the Remnant Security course interest and waitlist form.",
    )
    caption = qr_p.add_run("\nJoin the Remnant Security course waitlist")
    caption.font.name = "Avenir Next"
    caption.font.size = Pt(8)
    caption.font.bold = True
    caption.font.color.rgb = RGBColor.from_string(AMBER)
    move_paragraph_after(qr_p, cover_anchor)

    for table in doc.tables:
        # Pandoc's minimal reference document does not always contain Word's
        # built-in "Table Grid" style, so keep formatting self-contained.
        try:
            table.style = "Table Grid"
        except KeyError:
            pass
        table.autofit = True
        for row_index, row in enumerate(table.rows):
            keep_table_row_together(row)
            fill = NAVY if row_index == 0 else ("F3F8FA" if row_index % 2 else WHITE)
            for cell in row.cells:
                set_cell_shading(cell, fill)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs:
                        run.font.name = "Avenir Next"
                        run.font.size = Pt(8)
                        if row_index == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor.from_string(WHITE)

    core = doc.core_properties
    core.title = "How to Get Into Cybersecurity - 2026 Edition"
    core.author = "Ian Trimble | Remnant Security"
    core.subject = "Cybersecurity career field guide"
    core.keywords = "cybersecurity, careers, AI security, home lab, Security+, Remnant Security"

    doc.save(output_path)


def main() -> None:
    if len(sys.argv) != 5:
        raise SystemExit("usage: style_cyber_guide_docx.py INPUT OUTPUT LOGO QR")
    style_document(*(Path(arg).resolve() for arg in sys.argv[1:]))


if __name__ == "__main__":
    main()
